import json
import re
from pathlib import Path
from datetime import datetime

import openpyxl


# ---------------------------------------------------------
# CONFIG
# ---------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parent.parent

INPUT_FILE = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "Calendar_LHKH_online_2026.xlsx"
)

OUTPUT_FILE = (
    PROJECT_ROOT
    / "data"
    / "processed"
    / "jazzercise_corpus.json"
)

MONTHS_TO_USE = [
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
]


# ---------------------------------------------------------
# LOCATION NORMALIZATION
# ---------------------------------------------------------

def normalize_location(text):
    """
    Convert the different location spellings in the spreadsheet
    into consistent names.
    """

    lower = text.lower()

    if "wkk" in lower or "waikiki" in lower:
        return "Waikiki Community Center"

    if "st.peter" in lower or "st. peter" in lower or "st peter" in lower:
        return "St. Peters"

    if "palolo" in lower:
        return "Palolo Hongwanji"

    return None


# ---------------------------------------------------------
# TIME EXTRACTION
# ---------------------------------------------------------

def extract_time(text):
    """
    Extract times such as:
    8:30 am
    5pm
    6 pm
    6:00 pm
    """

    match = re.search(
        r"\b(\d{1,2})(?::(\d{2}))?\s*(am|pm)\b",
        text,
        re.IGNORECASE,
    )

    if not match:
        return None

    hour = int(match.group(1))
    minute = match.group(2) or "00"
    period = match.group(3).upper()

    return f"{hour}:{minute} {period}"


# ---------------------------------------------------------
# STATUS EXTRACTION
# ---------------------------------------------------------

def extract_status(text):

    if "no class" in text.lower():
        return "NO CLASS"

    return "Scheduled"


# ---------------------------------------------------------
# CLEAN CELL CONTENT
# ---------------------------------------------------------

def clean_lines(value):

    if value is None:
        return []

    return [
        line.strip(" -")
        for line in str(value).splitlines()
        if line.strip(" -")
    ]


# ---------------------------------------------------------
# CREATE LOCATION DOCUMENT
# ---------------------------------------------------------

def create_location_document(
    date_value,
    location,
    time,
    status,
    raw_entry,
    sheet_name,
    cell_reference,
):

    date_description = date_value.strftime(
        "%A, %B %d, %Y"
    )

    if status == "NO CLASS":

        page_content = (
            f"{location} has NO CLASS on "
            f"{date_description}."
        )

    elif time:

        page_content = (
            f"{location} has a Jazzercise class on "
            f"{date_description} at {time}."
        )

    else:

        page_content = (
            f"Jazzercise schedule information for "
            f"{location} on {date_description}: "
            f"{raw_entry}"
        )

    return {
        "page_content": page_content,

        "metadata": {
            "source": INPUT_FILE.name,
            "sheet": sheet_name,
            "date": date_value.strftime("%Y-%m-%d"),
            "day_of_week": date_value.strftime("%A"),
            "month": date_value.strftime("%B"),
            "year": date_value.year,
            "location": location,
            "time": time,
            "status": status,
            "raw_entry": raw_entry,
            "cell": cell_reference,
        },
    }


# ---------------------------------------------------------
# CREATE UNASSIGNED DOCUMENT
# ---------------------------------------------------------

def create_unassigned_document(
    date_value,
    raw_entry,
    sheet_name,
    cell_reference,
):
    """
    Preserve information such as a bare 'NO CLASS'
    without guessing which location it belongs to.
    """

    return {
        "page_content": (
            f"Schedule note for "
            f"{date_value.strftime('%A, %B %d, %Y')}: "
            f"{raw_entry}"
        ),

        "metadata": {
            "source": INPUT_FILE.name,
            "sheet": sheet_name,
            "date": date_value.strftime("%Y-%m-%d"),
            "day_of_week": date_value.strftime("%A"),
            "month": date_value.strftime("%B"),
            "year": date_value.year,
            "location": "Unspecified",
            "time": extract_time(raw_entry),
            "status": extract_status(raw_entry),
            "raw_entry": raw_entry,
            "cell": cell_reference,
        },
    }


# ---------------------------------------------------------
# MAIN CORPUS BUILD
# ---------------------------------------------------------

def build_corpus():

    print(f"Reading workbook: {INPUT_FILE}")

    workbook = openpyxl.load_workbook(
        INPUT_FILE,
        data_only=True
    )

    documents = []

    for sheet_name in MONTHS_TO_USE:

        worksheet = workbook[sheet_name]

        print(f"Processing {sheet_name}...")

        for row_number in range(
            1,
            worksheet.max_row
        ):

            for column_number in range(
                1,
                worksheet.max_column + 1
            ):

                date_cell = worksheet.cell(
                    row=row_number,
                    column=column_number,
                )

                if not isinstance(
                    date_cell.value,
                    datetime
                ):
                    continue

                # Ignore spillover dates from another month
                if (
                    date_cell.value.strftime("%B")
                    != sheet_name
                ):
                    continue

                schedule_cell = worksheet.cell(
                    row=row_number + 1,
                    column=column_number,
                )

                entries = clean_lines(
                    schedule_cell.value
                )

                for entry in entries:

                    location = normalize_location(
                        entry
                    )

                    if location:

                        document = (
                            create_location_document(
                                date_value=date_cell.value,
                                location=location,
                                time=extract_time(entry),
                                status=extract_status(entry),
                                raw_entry=entry,
                                sheet_name=sheet_name,
                                cell_reference=schedule_cell.coordinate,
                            )
                        )

                    else:

                        # Keep unusual notes without
                        # incorrectly assigning a location.
                        document = (
                            create_unassigned_document(
                                date_value=date_cell.value,
                                raw_entry=entry,
                                sheet_name=sheet_name,
                                cell_reference=schedule_cell.coordinate,
                            )
                        )

                    documents.append(document)

    OUTPUT_FILE.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    with open(
        OUTPUT_FILE,
        "w",
        encoding="utf-8",
    ) as file:

        json.dump(
            documents,
            file,
            indent=2,
            ensure_ascii=False,
        )

    print()
    print("---------------------------------------")
    print("Location-aware corpus complete")
    print("---------------------------------------")
    print(f"Documents created: {len(documents)}")
    print(f"Output: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_corpus()